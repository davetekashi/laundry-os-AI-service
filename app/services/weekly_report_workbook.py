from collections import defaultdict
from datetime import UTC, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


CONFIRMED_PAYMENT_STATUSES = {"confirmed", "completed", "paid", "success", "successful"}
NAVY = "123047"
BLUE = "159DD8"
PALE_BLUE = "EAF6FB"
PALE_GRAY = "F4F7F8"
WHITE = "FFFFFF"
GRID = "D9E3E8"
MONEY_FORMAT = '"NGN" #,##0.00;[Red]-"NGN" #,##0.00'
INTEGER_FORMAT = "#,##0"
PERCENT_FORMAT = "0.0%"


def _number(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _utc(value: datetime) -> datetime:
    if value.tzinfo is None:
        return value.replace(tzinfo=UTC)
    return value.astimezone(UTC)


def _payment_date(payment: dict) -> datetime | None:
    for field_name in ("paidAt", "transactionDate", "confirmedAt", "recordedAt", "createdAt"):
        value = payment.get(field_name)
        if isinstance(value, datetime):
            return _utc(value)
    return None


def _confirmed(payment: dict) -> bool:
    return str(payment.get("status") or "").casefold() in CONFIRMED_PAYMENT_STATUSES


def _signed_payment_amount(payment: dict) -> float:
    amount = _number(payment.get("amount"))
    return -amount if str(payment.get("transactionType") or "").casefold() == "refund" else amount


def _payment_method(payment: dict) -> str:
    raw = str(
        payment.get("offlineMethod")
        or payment.get("paymentChannel")
        or payment.get("method")
        or "Unspecified"
    )
    normalized = raw.replace("_", " ").strip().title()
    return {"Pos": "POS"}.get(normalized, normalized)


def _snapshot_name(document: dict) -> str:
    for field_name in ("customerSnapshot", "payerSnapshot"):
        snapshot = document.get(field_name) or {}
        if snapshot.get("fullName"):
            return str(snapshot["fullName"]).strip()
        name = " ".join(
            str(part).strip()
            for part in (snapshot.get("firstName"), snapshot.get("lastName"))
            if part
        )
        if name:
            return name
    return "Unknown customer"


def _customer_key(document: dict) -> str:
    for field_name in ("businessCustomerId", "laundryCustomerId", "userId"):
        value = document.get(field_name)
        if value is not None:
            return f"id:{value}"
    return f"name:{_snapshot_name(document).casefold()}"


def _physical_piece_count(order: dict) -> int:
    total_piece_count = order.get("totalPieceCount")
    if total_piece_count is not None:
        return max(int(_number(total_piece_count)), 0)

    nested_total = sum(
        max(int(_number(item.get("pieceCount"))), 0)
        for item in order.get("items") or []
        if item.get("pieceCount") is not None
    )
    if nested_total:
        return nested_total
    return max(int(_number(order.get("itemCount"))), 0)


def _expense_entries(expenses: list[dict]) -> list[dict]:
    rows: list[dict] = []
    for document in expenses:
        expense_date = document.get("expenseDate") or document.get("createdAt")
        entries = document.get("entries") or []
        if entries:
            for entry in entries:
                rows.append(
                    {
                        "date": expense_date,
                        "category": str(entry.get("category") or document.get("category") or "Uncategorized"),
                        "description": str(entry.get("subcategory") or entry.get("description") or ""),
                        "amount": _number(entry.get("amount")),
                    }
                )
            continue
        rows.append(
            {
                "date": expense_date,
                "category": str(document.get("category") or "Uncategorized"),
                "description": str(document.get("description") or document.get("title") or ""),
                "amount": _number(document.get("amount")) or _number(document.get("totalExpenses")),
            }
        )
    return rows


def build_weekly_workbook_data(
    raw_documents: dict,
    start_date: datetime,
    end_date: datetime,
) -> dict:
    start_date = _utc(start_date)
    end_date = _utc(end_date)
    selected_orders = raw_documents.get("orders_in_range") or []
    payment_orders = raw_documents.get("payment_orders") or []
    period_payments = [
        payment
        for payment in raw_documents.get("order_payments_in_range") or []
        if _confirmed(payment)
    ]
    selected_order_payments = [
        payment
        for payment in raw_documents.get("selected_order_payments") or []
        if _confirmed(payment)
        and (_payment_date(payment) is None or _payment_date(payment) <= end_date)
    ]

    orders_by_id = {
        order.get("_id"): order
        for order in selected_orders + payment_orders
        if order.get("_id") is not None
    }
    selected_order_ids = {
        order.get("_id") for order in selected_orders if order.get("_id") is not None
    }
    customers: dict[str, dict] = defaultdict(
        lambda: {
            "customer": "Unknown customer",
            "pieces": 0,
            "item_lines": 0,
            "order_value": 0.0,
            "received_for_weekly_orders": 0.0,
            "older_debt_recovered": 0.0,
            "outstanding_at_period_end": 0.0,
            "total_cash_received": 0.0,
        }
    )

    for order in selected_orders:
        key = _customer_key(order)
        row = customers[key]
        row["customer"] = _snapshot_name(order)
        row["pieces"] += _physical_piece_count(order)
        row["item_lines"] += max(int(_number(order.get("itemCount"))), 0)
        row["order_value"] += _number(order.get("totalPayable"))

    paid_through_end_by_order: dict[Any, float] = defaultdict(float)
    for payment in selected_order_payments:
        paid_through_end_by_order[payment.get("orderId")] += _signed_payment_amount(payment)
    for order in selected_orders:
        key = _customer_key(order)
        outstanding = max(
            _number(order.get("totalPayable"))
            - paid_through_end_by_order.get(order.get("_id"), 0.0),
            0.0,
        )
        customers[key]["outstanding_at_period_end"] += outstanding

    payment_methods: dict[str, float] = defaultdict(float)
    current_order_collections = 0.0
    older_debt_recovered = 0.0
    unallocated_collections = 0.0
    for payment in period_payments:
        amount = _signed_payment_amount(payment)
        payment_methods[_payment_method(payment)] += amount
        order = orders_by_id.get(payment.get("orderId"))
        if order is not None:
            key = _customer_key(order)
            row = customers[key]
            if row["customer"] == "Unknown customer":
                row["customer"] = _snapshot_name(payment)
            if payment.get("orderId") in selected_order_ids:
                row["received_for_weekly_orders"] += amount
                current_order_collections += amount
            else:
                order_date = order.get("createdAt")
                if isinstance(order_date, datetime) and _utc(order_date) < start_date:
                    row["older_debt_recovered"] += amount
                    older_debt_recovered += amount
                else:
                    unallocated_collections += amount
            row["total_cash_received"] += amount
        else:
            key = _customer_key(payment)
            row = customers[key]
            row["customer"] = _snapshot_name(payment)
            row["total_cash_received"] += amount
            unallocated_collections += amount

    customer_rows = sorted(
        customers.values(),
        key=lambda row: (row["order_value"], row["total_cash_received"]),
        reverse=True,
    )
    expense_rows = _expense_entries(raw_documents.get("expenses_in_range") or [])
    expense_categories: dict[str, float] = defaultdict(float)
    for row in expense_rows:
        expense_categories[row["category"]] += row["amount"]

    revenue = sum(_number(order.get("totalPayable")) for order in selected_orders)
    service_revenue = sum(_number(order.get("serviceTotal")) for order in selected_orders)
    logistics_revenue = sum(_number(order.get("logisticsTotal")) for order in selected_orders)
    total_collections = sum(_signed_payment_amount(payment) for payment in period_payments)
    expenses = sum(row["amount"] for row in expense_rows)
    outstanding = sum(row["outstanding_at_period_end"] for row in customer_rows)
    accounting_result = revenue - expenses
    cash_result = total_collections - expenses

    return {
        "metrics": {
            "order_count": len(selected_orders),
            "item_line_count": sum(max(int(_number(order.get("itemCount"))), 0) for order in selected_orders),
            "total_piece_count": sum(_physical_piece_count(order) for order in selected_orders),
            "revenue": revenue,
            "service_revenue": service_revenue,
            "logistics_revenue": logistics_revenue,
            "current_order_collections": current_order_collections,
            "older_debt_recovered": older_debt_recovered,
            "unallocated_collections": unallocated_collections,
            "total_collections": total_collections,
            "selected_order_outstanding_at_period_end": outstanding,
            "recorded_expenses": expenses,
            "accounting_result": accounting_result,
            "cash_result": cash_result,
            "profit_margin": accounting_result / revenue if revenue else 0.0,
        },
        "customer_rows": customer_rows,
        "payment_methods": dict(sorted(payment_methods.items())),
        "expense_rows": expense_rows,
        "expense_categories": dict(sorted(expense_categories.items())),
    }


def _section_header(sheet, row: int, start_column: int, end_column: int, title: str) -> None:
    sheet.merge_cells(
        start_row=row,
        start_column=start_column,
        end_row=row,
        end_column=end_column,
    )
    cell = sheet.cell(row, start_column, title)
    cell.font = Font(bold=True, color=WHITE, size=11)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center")


def _style_table(sheet, header_row: int, last_row: int, last_column: int) -> None:
    border = Border(
        left=Side(style="thin", color=GRID),
        right=Side(style="thin", color=GRID),
        top=Side(style="thin", color=GRID),
        bottom=Side(style="thin", color=GRID),
    )
    for cell in sheet[header_row]:
        if cell.column > last_column:
            break
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row in range(header_row + 1, last_row + 1):
        fill = PatternFill("solid", fgColor=PALE_GRAY if row % 2 == 0 else WHITE)
        for column in range(1, last_column + 1):
            cell = sheet.cell(row, column)
            cell.fill = fill
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def _autosize(sheet, maximum: int = 32) -> None:
    for column_cells in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column_cells) + 3, maximum)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = max(width, 12)


def render_weekly_summary_workbook(
    report_data: dict,
    business_name: str,
    start_date: datetime,
    end_date: datetime,
) -> bytes:
    workbook = Workbook()
    workbook.properties.title = f"{business_name} Weekly Business Summary"
    workbook.properties.creator = "Seanosis"
    summary = workbook.active
    summary.title = "Weekly Summary"
    summary.sheet_view.showGridLines = False
    metrics = report_data["metrics"]

    summary.merge_cells("A1:E1")
    summary["A1"] = f"{business_name} Weekly Business Summary"
    summary["A1"].font = Font(bold=True, size=20, color=NAVY)
    summary["A1"].alignment = Alignment(horizontal="center")
    summary.merge_cells("A2:E2")
    summary["A2"] = f"{_utc(start_date):%d %b %Y} to {_utc(end_date):%d %b %Y}"
    summary["A2"].font = Font(size=11, color="5B6B75")
    summary["A2"].alignment = Alignment(horizontal="center")

    _section_header(summary, 4, 1, 2, "Key Figures")
    key_figures = [
        ("Orders", metrics["order_count"], INTEGER_FORMAT),
        ("Physical Pieces", metrics["total_piece_count"], INTEGER_FORMAT),
        ("Item Lines", metrics["item_line_count"], INTEGER_FORMAT),
        ("Revenue", metrics["revenue"], MONEY_FORMAT),
        ("Service Revenue", metrics["service_revenue"], MONEY_FORMAT),
        ("Logistics Revenue", metrics["logistics_revenue"], MONEY_FORMAT),
        ("Collections for Weekly Orders", metrics["current_order_collections"], MONEY_FORMAT),
        ("Older Debt Recovered", metrics["older_debt_recovered"], MONEY_FORMAT),
        ("Unallocated Collections", metrics["unallocated_collections"], MONEY_FORMAT),
        ("Total Cash Received", metrics["total_collections"], MONEY_FORMAT),
        ("Selected-Order Outstanding", metrics["selected_order_outstanding_at_period_end"], MONEY_FORMAT),
        ("Recorded Expenses", metrics["recorded_expenses"], MONEY_FORMAT),
        ("Accounting Profit / Loss", metrics["accounting_result"], MONEY_FORMAT),
        ("Cash Surplus / Deficit", metrics["cash_result"], MONEY_FORMAT),
        ("Profit Margin", metrics["profit_margin"], PERCENT_FORMAT),
    ]
    border = Border(bottom=Side(style="thin", color=GRID))
    for row_number, (label, value, number_format) in enumerate(key_figures, start=5):
        summary.cell(row_number, 1, label).font = Font(bold=True, color="455A64")
        value_cell = summary.cell(row_number, 2, value)
        value_cell.font = Font(bold=True, color=NAVY)
        value_cell.number_format = number_format
        value_cell.alignment = Alignment(horizontal="right")
        summary.cell(row_number, 1).border = border
        value_cell.border = border

    _section_header(summary, 4, 4, 5, "Payment Channels")
    payment_row = 5
    for method, amount in report_data["payment_methods"].items():
        summary.cell(payment_row, 4, method)
        summary.cell(payment_row, 5, amount).number_format = MONEY_FORMAT
        summary.cell(payment_row, 5).alignment = Alignment(horizontal="right")
        payment_row += 1

    expense_start = max(payment_row + 2, 10)
    _section_header(summary, expense_start, 4, 5, "Expenses by Category")
    for row_number, (category, amount) in enumerate(
        report_data["expense_categories"].items(),
        start=expense_start + 1,
    ):
        summary.cell(row_number, 4, category)
        summary.cell(row_number, 5, amount).number_format = MONEY_FORMAT
        summary.cell(row_number, 5).alignment = Alignment(horizontal="right")

    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 22
    summary.column_dimensions["C"].width = 4
    summary.column_dimensions["D"].width = 28
    summary.column_dimensions["E"].width = 22
    summary.freeze_panes = "A4"
    summary.sheet_properties.pageSetUpPr.fitToPage = True
    summary.page_setup.orientation = "landscape"
    summary.page_setup.fitToWidth = 1
    summary.page_setup.fitToHeight = 1
    summary.print_area = f"A1:E{max(20, expense_start + len(report_data['expense_categories']))}"

    customer_sheet = workbook.create_sheet("Customer Summary")
    customer_sheet.sheet_view.showGridLines = False
    customer_headers = [
        "Customer",
        "Physical Pieces",
        "Item Lines",
        "Order Value",
        "Received for Weekly Orders",
        "Older Debt Recovered",
        "Outstanding at Period End",
        "Total Cash Received",
    ]
    customer_sheet.append(customer_headers)
    for row in report_data["customer_rows"]:
        customer_sheet.append(
            [
                row["customer"],
                row["pieces"],
                row["item_lines"],
                row["order_value"],
                row["received_for_weekly_orders"],
                row["older_debt_recovered"],
                row["outstanding_at_period_end"],
                row["total_cash_received"],
            ]
        )
    _style_table(customer_sheet, 1, max(customer_sheet.max_row, 1), len(customer_headers))
    for column in range(4, 9):
        for row in range(2, customer_sheet.max_row + 1):
            customer_sheet.cell(row, column).number_format = MONEY_FORMAT
    for column in (2, 3):
        for row in range(2, customer_sheet.max_row + 1):
            customer_sheet.cell(row, column).number_format = INTEGER_FORMAT
    customer_sheet.freeze_panes = "A2"
    customer_sheet.auto_filter.ref = customer_sheet.dimensions
    _autosize(customer_sheet)
    customer_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    customer_sheet.page_setup.orientation = "landscape"
    customer_sheet.page_setup.fitToWidth = 1
    customer_sheet.page_setup.fitToHeight = 0

    expense_sheet = workbook.create_sheet("Expenses")
    expense_sheet.sheet_view.showGridLines = False
    expense_headers = ["Date", "Category", "Description", "Amount"]
    expense_sheet.append(expense_headers)
    for row in report_data["expense_rows"]:
        date_value = row["date"]
        if isinstance(date_value, datetime) and date_value.tzinfo is not None:
            date_value = date_value.astimezone(UTC).replace(tzinfo=None)
        expense_sheet.append([date_value, row["category"], row["description"], row["amount"]])
    _style_table(expense_sheet, 1, max(expense_sheet.max_row, 1), len(expense_headers))
    for row in range(2, expense_sheet.max_row + 1):
        expense_sheet.cell(row, 1).number_format = "yyyy-mm-dd"
        expense_sheet.cell(row, 4).number_format = MONEY_FORMAT
    expense_sheet.freeze_panes = "A2"
    expense_sheet.auto_filter.ref = expense_sheet.dimensions
    _autosize(expense_sheet)
    expense_sheet.column_dimensions["A"].width = 16
    expense_sheet.column_dimensions["B"].width = 24
    expense_sheet.column_dimensions["C"].width = 52
    expense_sheet.column_dimensions["D"].width = 20
    expense_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    expense_sheet.page_setup.orientation = "landscape"
    expense_sheet.page_setup.fitToWidth = 1
    expense_sheet.page_setup.fitToHeight = 0

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
